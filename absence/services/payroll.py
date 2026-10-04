"""The monthly changes report for the bureau. Dates and units, never a
sickness category, never any pay arithmetic.

Which absence goes on which sheet is decided by the type's flags, not its
code. Only types with `payroll_reportable` appear at all; of those:

  Sickness       `health_sensitive` types: the name, the dates and whether it
                 was self-certified (Absence.self_certified), nothing about
                 the kind of illness (the category is never read here).
  Family leave   `is_family` types: the days in the period, expected and actual
                 dates, KIT days in the period.
  Unpaid         the rest that are not `paid` (the seeded UNPAID type and
                 dependants leave): the days in the period and the units
                 they cost (costing.cost_between).
  TOIL           not an absence sheet: the TOIL-earned and TOIL-taken ledger
                 lines of pot-backed types that are both `paid` and
                 `payroll_reportable`, so an unpaid TOIL type is on Unpaid
                 (its bookings) and not here. An approved claim's line is
                 reported in the month it was approved (the local date it
                 was written), with the day worked beside it, so a claim
                 approved after its month's report was made is still
                 reported once; any other line in the month of its date. A
                 lot carried forward at year end (noted
                 year_end.CARRIED_FROM, it keeps its earned date) is left
                 out: its original line is reported.

Starters, Leavers, Contract changes and Pay changes come from People. The
Leavers sheet also carries each leaver's balance in every pot-backed type
at the end of the leaving date, from the ledger lines dated up to it and its entitlement lines (a
negative one is a debt, a positive one leave unused), so both reach payroll.

`run` saves media/payroll/YYYY-MM.xlsx and records a PayrollRun. Generating
a month again records another run and replaces that month's file: the
report is rebuilt from the current data, so the file is always the latest
figures, and the runs (audited) show who generated it and when."""

import os
from calendar import monthrange
from datetime import date
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q, Sum
from django.utils import timezone
from openpyxl import Workbook

from absence.models import Absence, AbsenceType, LedgerEntry, PayrollRun
from absence.services import costing, pots, year_end
from people.models import Contract, Employment, PayRecord
from people.services import audit, contracts

SHEETS = ["Starters", "Leavers", "Contract changes", "Pay changes", "Sickness", "Unpaid", "Family leave", "TOIL"]

HEADERS = {
    "Starters": ["Name", "Start", "Contract type", "Weekly amount", "Unit",
                 "Account name", "Sort code", "Account number"],
    "Leavers": ["Name", "Last day", "Reason", "Unit"],      # then a balance per pot-backed type
    "Contract changes": ["Name", "From", "To", "Weekly amount", "Unit", "Basis", "Notes"],
    "Pay changes": ["Name", "From", "Basis", "Amount", "Reason"],
    "Sickness": ["Name", "From", "To", "Self-certified"],
    "Unpaid": ["Name", "From", "To", "Units", "Unit", "Type"],
    "Family leave": ["Name", "Type", "From", "To", "Expected start", "Actual start", "Expected return", "KIT days"],
    "TOIL": ["Name", "Date", "Units", "Kind", "Note", "Day worked", "Approved"],
}


def month(period):
    """(first day, last day) of a "YYYY-MM" month. ValueError when it is not one."""
    year, mon = (int(part) for part in period.split("-"))
    return date(year, mon, 1), date(year, mon, monthrange(year, mon)[1])


def _name(employment):
    return employment.employee.name


def _balance(employment, absence_type, day):
    """What the pot current on `day` held at the end of `day`: its ledger
    lines dated on or before it, plus its entitlement lines whatever their
    date (a leaver's entitlement is pro-rated by a revision line written
    when the end date is recorded, which can be after the last day). None when there is no such pot (not opened,
    or no contract or policy to open one from)."""
    try:
        pot = pots.lookup(employment, absence_type, day)
    except ValidationError:
        return None
    if pot is None:
        return None
    counted = Q(date__lte=day) | Q(kind__in=(LedgerEntry.Kind.ENTITLEMENT, LedgerEntry.Kind.REVISION))
    return float(pot.entries.filter(counted).aggregate(t=Sum("units"))["t"] or 0)


def _starters(start, end):
    for e in Employment.objects.filter(start_date__range=(start, end)).select_related("employee"):
        c = contracts.active_on(e, e.start_date).first()
        yield [_name(e), e.start_date, c.contract_type.name if c else "",
               float(contracts.contracted_amount(e, e.start_date)), c.contract_type.unit if c else "",
               e.employee.bank_account_name, e.employee.bank_sort_code, e.employee.bank_account_number]


def _leavers(start, end):
    types = list(AbsenceType.objects.filter(uses_pot=True))
    headers = HEADERS["Leavers"] + [f"{t.name} balance" for t in types]
    rows = []
    for e in Employment.objects.filter(end_date__range=(start, end)).select_related("employee"):
        rows.append([_name(e), e.end_date, e.get_leaving_reason_display(), contracts.unit(e, e.end_date) or ""]
                    + [_balance(e, t, e.end_date) for t in types])
    return headers, rows


def _absences(start, end):
    """The reportable absences of the period, sorted onto their sheets by
    flags. An absence spanning the period's edge is clipped to it, and an
    unpaid one costs only the days inside, so consecutive months never
    report the same day twice. The automatic bank-holiday rows are never
    reported: payroll pays bank holidays as part of the month."""
    sheets = {"Sickness": [], "Unpaid": [], "Family leave": []}
    live = (Absence.objects.filter(status=Absence.Status.APPROVED, start_date__lte=end, end_date__gte=start,
                                   absence_type__payroll_reportable=True, auto_bank_holiday=False)
            .select_related("employment__employee", "absence_type").prefetch_related("kit_days")
            .order_by("start_date", "id"))
    for a in live:
        t = a.absence_type
        first, last = max(a.start_date, start), min(a.end_date, end)
        if t.health_sensitive:
            sheets["Sickness"].append([_name(a.employment), first, last, "Yes" if a.self_certified else "No"])
        elif t.is_family:
            kit = sum(1 for k in a.kit_days.all() if first <= k.date <= last)
            sheets["Family leave"].append([_name(a.employment), t.name, first, last, a.expected_start,
                                           a.actual_start, a.expected_return, kit])
        elif not t.paid:
            units = costing.cost_between(a, start, end)
            sheets["Unpaid"].append([_name(a.employment), first, last, float(units),
                                     contracts.unit(a.employment, first) or "", t.name])
    return sheets


def _toil(start, end):
    """The TOIL sheet's rows, by the date each is reported under: an
    approved claim's line by the day it was approved (created_at, local),
    any other by its own date. Day worked is an earned line's date; Approved
    is filled for a claim's line only."""
    claimed = Q(toil_claim__isnull=False)
    lines = (LedgerEntry.objects.filter(
        kind__in=(LedgerEntry.Kind.TOIL_EARNED, LedgerEntry.Kind.TOIL_TAKEN),
        pot__absence_type__paid=True, pot__absence_type__payroll_reportable=True)
        .filter((claimed & Q(created_at__date__range=(start, end))) | (~claimed & Q(date__range=(start, end))))
        .exclude(kind=LedgerEntry.Kind.TOIL_EARNED, note__startswith=year_end.CARRIED_FROM)
        .select_related("pot__employment__employee", "toil_claim"))
    rows = []
    for line in lines:
        approved = timezone.localdate(line.created_at) if hasattr(line, "toil_claim") else None
        worked = line.date if line.kind == LedgerEntry.Kind.TOIL_EARNED else None
        rows.append(((approved or line.date), line.pk, [
            _name(line.pot.employment), approved or line.date, float(line.units), line.get_kind_display(),
            line.note, worked, approved]))
    return [row for _, _, row in sorted(rows, key=lambda r: (r[0], r[1]))]


def build(start, end):
    """The workbook for the period, one sheet per SHEETS entry, in order."""
    headers = dict(HEADERS)
    rows = {"Starters": list(_starters(start, end))}
    headers["Leavers"], rows["Leavers"] = _leavers(start, end)
    rows["Contract changes"] = [
        [_name(c.employment), c.from_date, c.to_date, float(c.weekly_amount), c.contract_type.unit,
         c.get_basis_display(), c.notes]
        for c in Contract.objects.filter(Q(from_date__range=(start, end)) | Q(to_date__range=(start, end)))
        .select_related("employment__employee", "contract_type").order_by("from_date", "id")]
    rows["Pay changes"] = [
        [_name(p.employment), p.from_date, p.get_basis_display(), float(p.amount), p.reason]
        for p in PayRecord.objects.filter(from_date__range=(start, end))
        .select_related("employment__employee").order_by("from_date", "id")]
    rows.update(_absences(start, end))
    rows["TOIL"] = _toil(start, end)
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(headers[name])
        for row in rows[name]:
            ws.append(row)
    return wb


def _save(wb, path):
    """Write beside the target and move into place, so a failed save never
    leaves a half-written report where the last good one was."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


@transaction.atomic
def run(actor, start, end):
    """Build the report for [start, end], save it as payroll/<start YYYY-MM>.xlsx
    under MEDIA_ROOT and record the run. The file is written last, so a
    failure earlier leaves no run and no file."""
    wb = build(start, end)
    counts = {ws.title: ws.max_row - 1 for ws in wb.worksheets}
    relative = Path("payroll") / f"{start:%Y-%m}.xlsx"
    row = PayrollRun.objects.create(period_start=start, period_end=end, generated_by=actor,
                                    path=relative.as_posix(), counts=counts)
    audit.record(actor, row, {"generated": ("", f"{start:%d %b %Y} to {end:%d %b %Y}")}, note=row.path)
    _save(wb, Path(settings.MEDIA_ROOT) / relative)
    return row
