from datetime import date
from decimal import Decimal

import openpyxl
from django.core.management import call_command

from absence.models import AbsenceType, LedgerEntry, PayrollRun
from absence.services import bookings, ledger, payroll, toil
from people.models import AuditEntry, PayRecord
from people.services import contracts, employments
from tests.factories import (absence_type, hours_employee, make_contract_type, make_employee, make_policy,
                             make_pot)

JUNE = (date(2026, 6, 1), date(2026, 6, 30))


def _sheet(wb, name):
    ws = wb[name]
    return [[c.value for c in row] for row in ws.iter_rows()]


def _cells(wb):
    return [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]


def test_sections(db, hr_admin, employee_user):
    starter = hours_employee(start=date(2026, 6, 15))
    leaver = hours_employee(employee=make_employee(first="Lee"), start=date(2025, 1, 1))
    employments.end(hr_admin, leaver, date(2026, 6, 20), "resigned")
    changed = hours_employee(employee=make_employee(first="Cha"), start=date(2025, 1, 1), amount=Decimal("20"))
    contracts.add(hr_admin, changed, make_contract_type(), Decimal("10"), date(2026, 6, 1))
    PayRecord.objects.create(employment=changed, from_date=date(2026, 6, 1), basis="annual", amount=26000)
    bookings.request(employee_user, changed, absence_type("SICK"), date(2026, 6, 3), date(2026, 6, 4), category="mental")
    unpaid = bookings.request(employee_user, changed, absence_type("UNPAID"), date(2026, 6, 10))
    bookings.approve(hr_admin, unpaid)
    ct = changed.contracts.first().contract_type
    make_policy(ct, "TOIL")
    toil.earn(hr_admin, changed, Decimal("2"), date(2026, 6, 5), "late")
    wb = payroll.build(*JUNE)
    assert wb.sheetnames == ["Starters", "Leavers", "Contract changes", "Pay changes", "Sickness", "Unpaid",
                             "Family leave", "TOIL"]
    assert _sheet(wb, "Starters")[1][0] == starter.employee.name
    assert _sheet(wb, "Leavers")[1][0] == leaver.employee.name
    assert any(r[0] == changed.employee.name for r in _sheet(wb, "Contract changes")[1:])
    assert _sheet(wb, "Pay changes")[1][3] == 26000
    sick_rows = _sheet(wb, "Sickness")
    assert sick_rows[1][1] == date(2026, 6, 3) and "mental" not in str(sick_rows)
    assert _sheet(wb, "Unpaid")[1][3] == 7.5
    assert _sheet(wb, "TOIL")[1][2] == 2


def test_sickness_category_appears_nowhere_in_the_workbook(db, employee_user):
    emp = hours_employee(start=date(2026, 6, 1))
    bookings.request(employee_user, emp, absence_type("SICK"), date(2026, 6, 3), category="mental")
    wb = payroll.build(*JUNE)
    assert _sheet(wb, "Sickness")[1][0] == emp.employee.name
    text = " ".join(_cells(wb)).lower()
    assert "mental" not in text and "category" not in text


def test_sheets_follow_type_flags_not_codes(db, employee_user, hr_admin):
    emp = hours_employee(start=date(2026, 6, 1))
    # a health-sensitive type of the practice's own goes to Sickness, dates only
    own = AbsenceType.objects.create(name="Long-term condition", code="LTC", payroll_reportable=True,
                                     health_sensitive=True, needs_approval=False)
    bookings.request(employee_user, emp, own, date(2026, 6, 3))
    # a type that is not payroll_reportable appears nowhere
    hidden = AbsenceType.objects.create(name="Sabbatical", code="SAB", paid=False, needs_approval=False)
    bookings.request(employee_user, emp, hidden, date(2026, 6, 10))
    # reportable and unpaid without being UNPAID (the seeded dependants leave) is still Unpaid
    bookings.request(employee_user, emp, absence_type("DEP"), date(2026, 6, 11))
    wb = payroll.build(*JUNE)
    assert [r[1] for r in _sheet(wb, "Sickness")[1:]] == [date(2026, 6, 3)]
    assert [r[4] for r in _sheet(wb, "Unpaid")[1:]] == ["Dependants leave"]
    assert "Sabbatical" not in " ".join(_cells(wb))


def test_family_leave_sheet_has_dates_and_kit_days(db, employee_user, hr_admin):
    emp = hours_employee(start=date(2025, 1, 1))
    a = bookings.request(employee_user, emp, absence_type("MAT"), date(2026, 5, 18), date(2027, 2, 12),
                         expected_start=date(2026, 5, 4), expected_return=date(2027, 2, 15))
    bookings.approve(hr_admin, a)
    bookings.add_kit_day(employee_user, a, date(2026, 8, 3))
    bookings.add_kit_day(employee_user, a, date(2026, 8, 4))
    rows = _sheet(payroll.build(*JUNE), "Family leave")
    assert rows[1] == [emp.employee.name, "Maternity leave", date(2026, 5, 4), None, date(2027, 2, 15), 2]


def test_toil_sheet_only_when_toil_is_paid(db, hr_admin):
    emp = hours_employee(start=date(2026, 6, 1))
    make_policy(emp.contracts.first().contract_type, "TOIL")
    toil.earn(hr_admin, emp, Decimal("2"), date(2026, 6, 5), "late")
    assert len(_sheet(payroll.build(*JUNE), "TOIL")) == 2
    AbsenceType.objects.filter(code="TOIL").update(paid=False)
    assert len(_sheet(payroll.build(*JUNE), "TOIL")) == 1        # headers only


def test_leavers_carry_pot_balances_at_the_leaving_date(db, hr_admin):
    emp = hours_employee(start=date(2025, 1, 1))
    pot = make_pot(emp, "AL", date(2026, 6, 1))
    ledger.write(pot, LedgerEntry.Kind.ADJUSTMENT, Decimal("-300"), hr_admin, note="over-taken", date=date(2026, 6, 2))
    employments.end(hr_admin, emp, date(2026, 6, 20), "resigned")
    rows = _sheet(payroll.build(*JUNE), "Leavers")
    header, row = rows[0], rows[1]
    assert header[:4] == ["Name", "Last day", "Reason", "Unit"]
    col = header.index("Annual leave balance")
    assert row[col] == float(ledger.balance(pot)) < 0           # a debt reaches payroll
    assert row[2] == "Resigned" and row[3] == "hours"
    assert row[header.index("Bank holiday balance")] in (None, 0)   # not a pot with a policy: blank or nil


def test_leaver_without_policy_for_a_type_has_a_blank_balance(db, hr_admin):
    emp = hours_employee(start=date(2025, 1, 1))
    employments.end(hr_admin, emp, date(2026, 6, 20), "resigned")
    rows = _sheet(payroll.build(*JUNE), "Leavers")
    assert rows[1][rows[0].index("Annual leave balance")] is None     # no pot open yet


def test_run_saves_file_and_row(db, hr_admin, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hours_employee(start=date(2026, 6, 15))
    run = payroll.run(hr_admin, date(2026, 6, 1), date(2026, 6, 30))
    assert (tmp_path / "payroll" / "2026-06.xlsx").exists()
    assert run.counts["Starters"] == 1 and PayrollRun.objects.count() == 1
    assert run.path == "payroll/2026-06.xlsx" and run.generated_by == hr_admin
    assert openpyxl.load_workbook(tmp_path / run.path)["Starters"].max_row == 2
    assert AuditEntry.objects.filter(model="absence.payrollrun", object_id=run.pk).exists()


def test_generating_a_period_again_records_a_second_run_and_replaces_the_file(db, hr_admin, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    hours_employee(start=date(2026, 6, 15))
    first = payroll.run(hr_admin, *JUNE)
    hours_employee(employee=make_employee(first="Late"), start=date(2026, 6, 20))
    second = payroll.run(hr_admin, *JUNE)
    assert PayrollRun.objects.count() == 2 and first.counts["Starters"] == 1 and second.counts["Starters"] == 2
    assert list((tmp_path / "payroll").iterdir()) == [tmp_path / "payroll" / "2026-06.xlsx"]
    assert openpyxl.load_workbook(tmp_path / second.path)["Starters"].max_row == 3


def test_view_requires_hr_admin(employee_client, admin_client, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    assert employee_client.get("/absence/payroll/").status_code == 403
    assert employee_client.post("/absence/payroll/", {"period": "2026-06"}).status_code == 403
    assert admin_client.get("/absence/payroll/").status_code == 200
    r = admin_client.post("/absence/payroll/", {"period": "2026-06"})
    assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
    assert "2026-06.xlsx" in r["Content-Disposition"]
    assert PayrollRun.objects.count() == 1


def test_view_lists_runs_and_rejects_a_bad_period(admin_client, hr_admin, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    payroll.run(hr_admin, *JUNE)
    assert b"June 2026" in admin_client.get("/absence/payroll/").content
    r = admin_client.post("/absence/payroll/", {"period": "2026-13"})
    assert r.status_code == 200 and not r["Content-Type"].startswith("application/vnd") and PayrollRun.objects.count() == 1
    assert admin_client.post("/absence/payroll/", {"period": "junk"}).status_code == 200


def test_command_writes_the_month(db, settings, tmp_path, capsys):
    settings.MEDIA_ROOT = tmp_path
    hours_employee(start=date(2026, 2, 10))
    call_command("payroll_report", period="2026-02")
    run = PayrollRun.objects.get()
    assert (run.period_start, run.period_end) == (date(2026, 2, 1), date(2026, 2, 28)) and run.generated_by is None
    assert str(tmp_path / "payroll" / "2026-02.xlsx") in capsys.readouterr().out
