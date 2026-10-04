import pytest
from django.core.exceptions import ValidationError

from people.models import AuditEntry
from people.services import employees
from tests.factories import make_employee

pytestmark = pytest.mark.django_db


def test_bank_details_are_editable_through_the_service_and_audited(hr_admin):
    e = make_employee()
    employees.update(hr_admin, e, bank_account_name="S Patel", bank_sort_code="12-34-56",
                     bank_account_number="12345678")
    e.refresh_from_db()
    assert (e.bank_sort_code, e.bank_account_number) == ("12-34-56", "12345678")
    fields = set(AuditEntry.objects.filter(object_id=e.pk).values_list("field", flat=True))
    assert {"bank_account_name", "bank_sort_code", "bank_account_number"} <= fields


@pytest.mark.parametrize("field,value", [("bank_sort_code", "123456"), ("bank_sort_code", "12-34-5x"),
                                         ("bank_account_number", "1234567"), ("bank_account_number", "1234567a")])
def test_malformed_bank_details_are_refused(hr_admin, field, value):
    e = make_employee()
    with pytest.raises(ValidationError):
        employees.update(hr_admin, e, **{field: value})


def test_opening_an_employee_with_bank_details_is_audited_as_viewed(admin_client, hr_admin):
    e = make_employee(bank_account_number="12345678", bank_sort_code="12-34-56")
    admin_client.get(f"/admin/people/employee/{e.pk}/change/")
    assert AuditEntry.objects.filter(object_id=e.pk, kind="viewed", field="bank").exists()


def test_the_payroll_starters_sheet_carries_bank_details(hr_admin, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    from openpyxl import load_workbook

    from absence.services import payroll
    from tests.factories import hours_employee
    emp = hours_employee()
    employees.update(hr_admin, emp.employee, bank_account_name="S Patel", bank_sort_code="12-34-56",
                     bank_account_number="12345678")
    start, end = payroll.month(f"{emp.start_date:%Y-%m}")
    run = payroll.run(hr_admin, start, end)
    ws = load_workbook(tmp_path / run.path)["Starters"]
    headers = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert headers[-3:] == ["Account name", "Sort code", "Account number"]
    assert row[-3:] == ["S Patel", "12-34-56", "12345678"]
